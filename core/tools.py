import os
import re
import time
import datetime
import threading
import webbrowser
import urllib.parse
import mss
from comtypes import CLSCTX_ALL
from pycaw.pycaw import AudioUtilities, IAudioEndpointVolume

class SystemTools:
    # --- CONTACT BOOK ---
    CONTACTS = {
        "mom": "+919876543210",
        "dad": "+919876543211",
        "alex": "+919876543212",
        "bimbo": "+917993328366"
    }

    DEFAULT_COUNTRY_CODE = "+91"

    # --- HELPER: RESOLVE RECIPIENT TO PHONE NUMBER ---
    @classmethod
    def _resolve_number(cls, recipient: str) -> str:
        clean_name = recipient.lower().strip()
        if clean_name in cls.CONTACTS:
            return cls.CONTACTS[clean_name]
        
        digits = re.sub(r'\D', '', recipient)
        if len(digits) == 10:
            return f"{cls.DEFAULT_COUNTRY_CODE}{digits}"
        elif len(digits) > 10:
            return f"+{digits}"
        return None

    # --- QUICK RELATIVE TIMERS & REMINDERS ---
    @staticmethod
    def set_relative_timer(user_text: str, tts_engine=None) -> str:
        try:
            numbers = re.findall(r'\d+', user_text)
            if not numbers:
                return "Please specify a time duration in minutes, seconds, or hours, Boss."
            
            duration = int(numbers[0])
            seconds = duration
            
            if "minute" in user_text:
                seconds = duration * 60
            elif "hour" in user_text:
                seconds = duration * 3600

            # Extract task description
            clean_text = user_text.lower().replace("remind me", "").replace("in ", "").replace("after ", "")
            clean_text = re.sub(r'\d+\s*(minute|second|hour)s?', '', clean_text).replace("to ", "").strip()
            
            reminder_msg = clean_text if clean_text else "resume your task"

            def alert_user():
                msg = f"Reminder Boss: Time to {reminder_msg}."
                print(f"\n[⏰ REMINDER ALERT]: {msg}\n")
                if tts_engine:
                    tts_engine.speak(msg)

            # Fire non-blocking background thread
            timer_thread = threading.Timer(seconds, alert_user)
            timer_thread.daemon = True
            timer_thread.start()

            unit = "minute" if "minute" in user_text else ("hour" if "hour" in user_text else "second")
            if duration > 1:
                unit += "s"

            return f"Timer set for {duration} {unit}. I will alert you out loud when time is up, Boss."

        except Exception as e:
            return f"Failed to set reminder timer, Boss: {e}"

    # --- WHATSAPP DESKTOP MESSAGING ---
    @classmethod
    def send_whatsapp_message(cls, recipient: str, message: str) -> str:
        try:
            import pyautogui
            import subprocess

            phone_number = cls._resolve_number(recipient)
            if not phone_number:
                return f"Invalid phone number or contact '{recipient}', Boss."

            clean_number = phone_number.replace("+", "").replace(" ", "").replace("-", "")
            encoded_message = urllib.parse.quote(message)
            
            desktop_cmd = f'start whatsapp://send?phone={clean_number}^&text={encoded_message}'
            subprocess.run(["cmd", "/c", desktop_cmd], shell=True)
            
            time.sleep(4)
            
            screen_width, screen_height = pyautogui.size()
            pyautogui.click(screen_width // 2, screen_height // 2)
            time.sleep(0.5)
            
            pyautogui.press("enter")
            
            contact_display = recipient.capitalize() if recipient in cls.CONTACTS else phone_number
            return f"WhatsApp message sent to {contact_display}, Boss."
        except Exception as e:
            return f"Failed to send WhatsApp message, Boss: {e}"

    # --- YOUTUBE AUTOMATION ---
    @staticmethod
    def play_youtube_song(user_text: str) -> str:
        try:
            import pywhatkit
            # Clean sentence-level filler without altering words like 'songs' or 'playlist'
            clean_query = re.sub(r'^(open youtube and|play a song|play songs|play song|play)\s*', '', user_text, flags=re.IGNORECASE)
            clean_query = re.sub(r'\s*(on youtube|from youtube)$', '', clean_query, flags=re.IGNORECASE).strip()
            
            if not clean_query:
                clean_query = "trending music"
                
            pywhatkit.playonyt(clean_query)
            return f"Playing '{clean_query}' on YouTube, Boss."
        except Exception as e:
            return f"Failed to play on YouTube, Boss: {e}"

    # --- AUDIO & VOLUME CONTROL ---
    @staticmethod
    def set_volume(level: int) -> str:
        try:
            level = max(0, min(100, level))
            devices = AudioUtilities.GetSpeakers()
            interface = devices.Activate(IAudioEndpointVolume._iid_, CLSCTX_ALL, None)
            volume = interface.QueryInterface(IAudioEndpointVolume)
            volume.SetMasterVolumeLevelScalar(level / 100.0, None)
            return f"Volume set to {level} percent, Boss."
        except Exception as e:
            return f"Failed to set volume, Boss: {e}"

    @staticmethod
    def mute_audio(mute: bool = True) -> str:
        try:
            devices = AudioUtilities.GetSpeakers()
            interface = devices.Activate(IAudioEndpointVolume._iid_, CLSCTX_ALL, None)
            volume = interface.QueryInterface(IAudioEndpointVolume)
            volume.SetMute(int(mute), None)
            status = "muted" if mute else "unmuted"
            return f"Audio {status}, Boss."
        except Exception as e:
            return f"Failed to toggle mute, Boss: {e}"

    # --- SCREENSHOT MANAGEMENT ---
    @staticmethod
    def take_screenshot() -> str:
        try:
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"screenshot_{timestamp}.png"
            target_dir = os.path.join(os.path.expanduser("~"), "Pictures", "Screenshots")
            os.makedirs(target_dir, exist_ok=True)
            
            filepath = os.path.join(target_dir, filename)
            with mss.mss() as sct:
                sct.shot(output=filepath)

            return "Screenshot saved to your Screenshots folder, Boss."
        except Exception as e:
            return f"Failed to capture screenshot, Boss: {e}"

    # --- UNIVERSAL MICROSOFT WORD DOCUMENT GENERATOR ---
    @staticmethod
    def create_word_document(user_text: str, assistant_engine=None) -> str:
        try:
            import docx
            from docx.shared import Pt, Inches

            # 1. Determine Target Save Directory from User Command
            user_low = user_text.lower()
            home_dir = os.path.expanduser("~")
            
            if "on desktop" in user_low or "to desktop" in user_low or "in desktop" in user_low:
                target_dir = os.path.join(home_dir, "Desktop")
            elif "in downloads" in user_low or "to downloads" in user_low:
                target_dir = os.path.join(home_dir, "Downloads")
            else:
                path_match = re.search(r'in\s+([a-zA-Z]:\\[^\s]+)', user_text)
                if path_match:
                    target_dir = path_match.group(1)
                else:
                    target_dir = os.path.join(home_dir, "Documents")

            os.makedirs(target_dir, exist_ok=True)

            # 2. Generate Content using Local LLM Engine
            # Inside create_word_document in core/tools.py:
            if assistant_engine:
                prompt = (
                    f"You are a professional document writer. Generate the FULL, COMPLETE content for this request: '{user_text}'.\n"
                    f"Requirements:\n"
                    f"- Write out all paragraphs, details, dates, and placeholders completely.\n"
                    f"- Do NOT summarize or use shortcuts.\n"
                    f"- Output ONLY the body text of the document. Do not include any intro like 'Here is your letter' or extra commentary."
                )
                generated_text = assistant_engine.process_message(prompt)
                clean_text = re.sub(r'```.*?\n|```', '', generated_text).strip()
                
                
            else:
                clean_text = f"Document content for request:\n\n{user_text}"

            # 3. Create and Format Word Document (.docx)
            doc = docx.Document()
            
            for section in doc.sections:
                section.top_margin = Inches(1)
                section.bottom_margin = Inches(1)
                section.left_margin = Inches(1)
                section.right_margin = Inches(1)

            for line in clean_text.split('\n'):
                line_str = line.strip()
                if not line_str:
                    continue
                
                if line_str.isupper() and len(line_str) < 50:
                    p = doc.add_paragraph()
                    run = p.add_run(line_str)
                    run.font.name = 'Calibri'
                    run.font.size = Pt(14)
                    run.bold = True
                else:
                    p = doc.add_paragraph()
                    run = p.add_run(line_str)
                    run.font.name = 'Calibri'
                    run.font.size = Pt(11)

            # 4. Generate Safe Filename & Save
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"NOVA_Doc_{timestamp}.docx"
            filepath = os.path.join(target_dir, filename)
            doc.save(filepath)

            # 5. Open File directly in Microsoft Word
            os.system(f'start winword "{filepath}"')

            folder_name = os.path.basename(target_dir) if os.path.basename(target_dir) else target_dir
            return f"Generated document saved to {folder_name} and opened in Microsoft Word, Boss."

        except Exception as e:
            return f"Failed to generate Word document, Boss: {e}"

    # --- WORKSPACE & APPLICATION CONTROL ---
    @staticmethod
    def open_app(app_name: str) -> str:
        name_lower = app_name.lower()

        if "opencode" in name_lower:
            opencode_path = r"C:\Users\Prasanna Kumar\AppData\Local\Programs\opencode\OpenCode.exe"
            if os.path.exists(opencode_path):
                os.startfile(opencode_path)
            else:
                os.system("start opencode")
            return "Opening OpenCode Application, Boss."

        app_map = [
            ("vs code", "code"),
            ("vscode", "code"),
            ("code", "code"),
            ("chrome", "chrome"),
            ("notepad", "notepad"),
            ("calculator", "calc"),
            ("cmd", "start cmd"),
            ("terminal", "start wt"),
            ("explorer", "explorer"),
            ("whatsapp", "whatsapp:")
        ]
        
        for key, command in app_map:
            if key in name_lower:
                os.system(f"start {command}")
                return f"Opening {key.capitalize()}, Boss."
        return f"Could not locate application mapping for '{app_name}', Boss."

    # --- COMMAND ROUTER ---
    @classmethod
    def process_command(cls, user_text: str, tts_engine=None, assistant_engine=None):
        text = user_text.lower().strip()

        # 1. QUICK RELATIVE REMINDERS & TIMERS
        if "remind me" in text and ("after" in text or "in" in text or "minute" in text or "second" in text or "hour" in text):
            response = cls.set_relative_timer(user_text, tts_engine)
            return True, response

        # 2. UNIVERSAL WORD DOCUMENT GENERATION
        if ("word" in text or "doc" in text or "leave letter" in text or "report" in text or "essay" in text) and ("write" in text or "draft" in text or "create" in text or "make" in text or "type" in text or "open" in text):
            return True, cls.create_word_document(user_text, assistant_engine=assistant_engine)

        # 3. WHATSAPP MESSAGING
        if "whatsapp" in text and ("send" in text or "message" in text or "text" in text):
            message_match = re.search(r"'(.*?)'|\"(.*?)\"", user_text)
            if message_match:
                message_content = message_match.group(1) or message_match.group(2)
            elif "saying" in text:
                message_content = text.split("saying")[-1].strip()
            else:
                message_content = "Hello"

            recipient = None
            for name in cls.CONTACTS:
                if name in text:
                    recipient = name
                    break
            
            if not recipient:
                digits = re.findall(r'\d+', text)
                recipient = digits[0] if digits else "unknown"

            response = cls.send_whatsapp_message(recipient, message_content)
            return True, response

        # 4. YOUTUBE AUTOMATION
        if "play" in text and ("song" in text or "playlist" in text or "youtube" in text or "from" in text):
            return True, cls.play_youtube_song(user_text)

        # 5. AUDIO & VOLUME CONTROL
        if "set volume to" in text or "volume to" in text:
            numbers = re.findall(r'\d+', text)
            if numbers:
                val = int(numbers[0])
                return True, cls.set_volume(val)

        if "mute audio" in text or "mute volume" in text or "mute" in text:
            return True, cls.mute_audio(True)

        if "unmute" in text or "unmute audio" in text:
            return True, cls.mute_audio(False)

        # 6. SCREENSHOTS
        if "screenshot" in text or "take a shot" in text:
            return True, cls.take_screenshot()

        # 7. GENERIC APP LAUNCHING
        if "open" in text or "launch" in text:
            return True, cls.open_app(user_text)

        return False, None