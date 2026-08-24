import os
from pathlib import Path
from typing import Optional
from rag.document_models import Document


class DocumentLoader:
    """
    Offline local document reader supporting .txt, .md, .pdf, and .docx formats.
    """

    SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}

    @classmethod
    def load_document(cls, file_path_str: str) -> Document:
        path = Path(file_path_str).resolve()

        if not path.exists():
            raise FileNotFoundError(f"Document file not found at: {file_path_str}")

        ext = path.suffix.lower()
        if ext not in cls.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported document extension '{ext}'. Supported: {cls.SUPPORTED_EXTENSIONS}")

        content = ""
        metadata = {"file_size_bytes": path.stat().st_size}

        try:
            if ext in [".txt", ".md"]:
                content = cls._load_text(path)
            elif ext == ".pdf":
                content, pdf_meta = cls._load_pdf(path)
                metadata.update(pdf_meta)
            elif ext == ".docx":
                content = cls._load_docx(path)
        except Exception as e:
            raise RuntimeError(f"Failed to read document {path.name}: {str(e)}")

        if not content.strip():
            raise ValueError(f"Document '{path.name}' contains no readable text.")

        return Document(
            filename=path.name,
            source_path=str(path),
            file_type=ext.lstrip("."),
            content=content,
            metadata=metadata,
        )

    @staticmethod
    def _load_text(path: Path) -> str:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    @staticmethod
    def _load_pdf(path: Path) -> tuple[str, dict]:
        import pypdf
        reader = pypdf.PdfReader(str(path))
        text_parts = []
        for idx, page in enumerate(reader.pages):
            extracted = page.extract_text()
            if extracted:
                text_parts.append(f"[Page {idx + 1}]\n{extracted}")
        return "\n\n".join(text_parts), {"total_pages": len(reader.pages)}

    @staticmethod
    def _load_docx(path: Path) -> str:
        import docx
        doc = docx.Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)