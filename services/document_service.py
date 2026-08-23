import re
from pathlib import Path
from typing import Optional, List, Dict, Any
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


class DocumentService:
    """
    Physical .docx creation and formatting service for NOVA V2.
    Handles document structural layout, styles, margins, and safe user file persistence.
    """

    DEFAULT_OUTPUT_DIR = Path.home() / "Documents" / "NOVA"

    @classmethod
    def get_output_dir(cls) -> Path:
        output_dir = cls.DEFAULT_OUTPUT_DIR
        output_dir.mkdir(parents=True, exist_ok=True)
        return output_dir

    @classmethod
    def create_document(
        cls,
        title: str,
        content_blocks: List[Dict[str, Any]],
        document_type: str = "report",
        filename: Optional[str] = None,
    ) -> Path:
        """
        Creates and formats a .docx file on disk.

        content_blocks expect items like:
        [
            {"type": "heading", "text": "Executive Summary", "level": 1},
            {"type": "paragraph", "text": "This report details..."},
            {"type": "bullet", "text": "Point A"},
            {"type": "numbered", "text": "Step 1"}
        ]
        """
        doc = docx.Document()

        # Set 1-inch margins across all sections
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Title Header
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(title)
        title_run.font.name = "Calibri"
        title_run.font.size = Pt(22)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(31, 78, 121)  # Deep Navy Blue

        doc.add_paragraph()  # Spacing

        # Render Content Blocks
        for block in content_blocks:
            b_type = block.get("type", "paragraph")
            b_text = block.get("text", "").strip()
            if not b_text:
                continue

            if b_type == "heading":
                level = block.get("level", 1)
                h_p = doc.add_paragraph()
                h_p.paragraph_format.space_before = Pt(12)
                h_p.paragraph_format.space_after = Pt(4)
                h_run = h_p.add_run(b_text)
                h_run.font.name = "Calibri"
                h_run.font.bold = True
                if level == 1:
                    h_run.font.size = Pt(16)
                    h_run.font.color.rgb = RGBColor(31, 78, 121)
                elif level == 2:
                    h_run.font.size = Pt(13)
                    h_run.font.color.rgb = RGBColor(89, 89, 89)
                else:
                    h_run.font.size = Pt(11)
                    h_run.font.color.rgb = RGBColor(38, 38, 38)

            elif b_type == "bullet":
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                run = p.add_run(b_text)
                run.font.name = "Calibri"
                run.font.size = Pt(11)

            elif b_type == "numbered":
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(3)
                run = p.add_run(b_text)
                run.font.name = "Calibri"
                run.font.size = Pt(11)

            else:  # Normal paragraph
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                run = p.add_run(b_text)
                run.font.name = "Calibri"
                run.font.size = Pt(11)

        # Generate Safe Filename
        if not filename:
            clean_title = re.sub(r"[^\w\s-]", "", title).strip().replace(" ", "_")
            if not clean_title:
                clean_title = f"{document_type}_document"
            filename = f"{clean_title}.docx"
        elif not filename.endswith(".docx"):
            filename = f"{filename}.docx"

        output_path = cls.get_output_dir() / filename
        doc.save(str(output_path))
        return output_path